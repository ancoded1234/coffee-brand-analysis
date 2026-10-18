"""Editable editorial PowerPoint and presenter notes, generated from reviewed CSVs."""
from pathlib import Path
import json,html
import pandas as pd
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION,XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
R=Path(__file__).resolve().parents[1];P=R/'presentation';A=P/'assets'
B=pd.read_csv(R/'data/comparison.csv').fillna('미공개');S=pd.read_csv(R/'data/sources.csv').fillna('미공개');M=json.loads((R/'analysis/metrics.json').read_text(encoding='utf-8'))
CREAM='F4EFE6';INK='30231D';BROWN='754B36';GREEN='63765D';BEIGE='E8DFD1';MUTED='807164';WHITE='FFFCF7';LINE='CFC1AE';ACCENT='B48557'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='비싼 커피는 정말 더 좋은 원두를 사용할까?';prs.core_properties.subject='한국 6개 카페 프랜차이즈 공개 자료 분석';prs.core_properties.author='Coffee Editorial / Data Literacy';prs.core_properties.keywords='실제 웹 수집, Firecrawl, 공개 정보, HOT 아메리카노'
slides=[];notes=[]
def rgb(c):return RGBColor.from_string(c)
def rect(sl,x,y,w,h,fill=BEIGE,line=None,radius=False,name=''):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background() if line is None else None
    if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(.6)
    if radius:
        try:sh.adjustments[0]=.08
        except Exception:pass
    if name:sh.name=name
    return sh
def text(sl,s,x,y,w,h=0.5,size=18,color=INK,bold=False,font='맑은 고딕',align=None,name=''):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,line in enumerate(str(s).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(5);p.space_before=Pt(0)
        if align is not None:p.alignment=align
    if name:sh.name=name
    return sh
def line(sl,x,y,w,color=LINE):return rect(sl,x,y,w,.012,color)
def base(n,kicker,title,sub='',dark=False,source=''):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(INK if dark else CREAM)
    fg=CREAM if dark else INK
    text(sl,'COFFEE / EVIDENCE / VALUE',.55,.29,6,.3,9,MUTED if not dark else 'C3AE98',font='Arial',bold=True)
    text(sl,kicker.upper(),.55,.85,11.9,.35,11,GREEN if not dark else 'C3AE98',bold=True,font='Arial')
    text(sl,title,.55,1.28,12.1,1.18,32,fg,True)
    if sub:text(sl,sub,.58,2.35,12.0,.65,14,MUTED if not dark else 'C3AE98')
    line(sl,.55,6.94,12.2,LINE if not dark else '685244')
    text(sl,source or '수집 2026.10.18 · 공개 자료 조사 · 원두 품질 순위와 인과관계를 산출하지 않음',.55,7.08,11.5,.24,8.5,MUTED if not dark else 'C3AE98')
    text(sl,f'{n:02} / 13',12.0,7.07,.78,.3,9,fg,True,font='Arial',align=PP_ALIGN.RIGHT)
    slides.append(dict(number=n,kicker=kicker,title=title,subtitle=sub,source=source,dark=dark))
    return sl
def note(sl,s,ids=''):
    refs=[]
    for sid in ids.split('|'):
        if not sid:continue
        row=S[S['출처ID']==sid].iloc[0];refs.append(f"[{sid}] {row['출처종류']} | 게시/적용: {row['게시일']} | 수집 {row['수집일']}\n{row['출처URL']}")
    full=s+'\n\n'+ '\n'.join(refs);sl.notes_slide.notes_text_frame.text=full;notes.append(full);slides[-1]['notes']=s;slides[-1]['source_ids']=ids.split('|') if ids else []
def card(sl,x,y,w,h,label,headline,body,fill=WHITE):
    rect(sl,x,y,w,h,fill,radius=True);text(sl,label,x+.22,y+.18,w-.44,.3,10,GREEN,True,font='Arial');text(sl,headline,x+.22,y+.67,w-.44,.9,23,INK,True);text(sl,body,x+.22,y+1.65,w-.44,h-1.8,15,MUTED)
def chart(sl,cats,vals,x,y,w,h,title='',pointcolors=None,maxval=None):
    data=CategoryChartData();data.categories=cats;data.add_series(title or '값',vals)
    ch=sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    ch.has_legend=False;ch.chart_style=10
    ch.chart_title.has_text_frame=True if title else False
    ch.has_title=bool(title)
    if title:ch.chart_title.text_frame.text=title
    ch.category_axis.reverse_order=True;ch.category_axis.tick_labels.font.name='맑은 고딕';ch.category_axis.tick_labels.font.size=Pt(12);ch.category_axis.tick_labels.font.color.rgb=rgb(INK)
    ch.value_axis.tick_labels.font.name='Arial';ch.value_axis.tick_labels.font.size=Pt(10);ch.value_axis.minimum_scale=0
    if maxval:ch.value_axis.maximum_scale=maxval
    ch.value_axis.has_major_gridlines=True;ch.value_axis.major_gridlines.format.line.color.rgb=rgb(LINE);ch.value_axis.major_gridlines.format.line.width=Pt(.4)
    ch.category_axis.major_tick_mark=XL_TICK_MARK.NONE;ch.value_axis.major_tick_mark=XL_TICK_MARK.NONE
    ch.category_axis.format.line.fill.background();ch.value_axis.format.line.fill.background()
    plot=ch.plots[0];plot.gap_width=65;plot.has_data_labels=True;plot.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;plot.data_labels.font.name='Arial';plot.data_labels.font.size=Pt(13);plot.data_labels.font.bold=True;plot.data_labels.font.color.rgb=rgb(INK);plot.data_labels.number_format='#,##0'
    for i,p in enumerate(ch.series[0].points):p.format.fill.solid();p.format.fill.fore_color.rgb=rgb((pointcolors or [BROWN]*len(vals))[i]);p.format.line.fill.background()
    return ch
def grid(sl,headers,rows,widths,x=.6,y=3.05,rowh=.45,size=12):
    xx=x
    for h,w in zip(headers,widths):text(sl,h,xx+.07,y,w-.14,.4,10,GREEN,True);xx+=w
    yy=y+.5
    for r,row in enumerate(rows):
        rect(sl,x,yy,sum(widths),rowh-.04,WHITE if r%2==0 else BEIGE,radius=True);xx=x
        for j,(v,w) in enumerate(zip(row,widths)):
            text(sl,v,xx+.10,yy+.10,w-.20,rowh-.14,size,INK,j==0);xx+=w
        yy+=rowh

sl=base(1,'A coffee editorial / Data literacy','',dark=False)
sl.shapes.add_picture(str(A/'coffee-editorial-flow.jpg'),Inches(0),Inches(0),width=prs.slide_width,height=prs.slide_height)
rect(sl,0,0,6.7,7.5,CREAM)
text(sl,'THE PRICE OF A CUP',.58,.55,6,.4,12,GREEN,True,font='Arial')
text(sl,'비싼 커피는\n정말 더 좋은\n원두를 사용할까?',.55,1.50,6.3,3.2,37,INK,True)
text(sl,'한국 6대 카페 프랜차이즈의\n원두 구성과 판매가격 데이터 분석',.6,5.15,5.8,.9,17,MUTED)
text(sl,'6 BRANDS     /     PUBLIC EVIDENCE     /     2026.10.18',.6,6.81,6.2,.3,9,INK,True,font='Arial')
text(sl,'Google Flow AI 생성 이미지 · 실제 브랜드·품질의 증거 아님',7.05,7.1,5.5,.25,8,INK)
slides[-1]['title']='비싼 커피는 정말 더 좋은 원두를 사용할까?'
note(sl,'오늘의 질문은 비싼 커피에 지불한 돈이 더 좋은 원두의 대가인지입니다. 먼저 가격과 원두의 공개 정보를 모으고, 무엇을 확인할 수 있는지부터 따져 봤습니다. 표지 이미지는 Google Flow로 만든 설명용 비주얼이며, 브랜드 제품이나 원두 품질의 증거가 아닙니다.')

sl=base(2,'01 / Research question','가격표는 보인다. 원두의 품질도 보일까?','가설: 더 비싼 커피에는 공개 자료로 구별할 수 있는 원두 특성이 있을 것이다.')
for x,lab,head,body in [(0.6,'OBSERVE','판매가격','HOT 기본 아메리카노의 공개 가격을 수집한다.'),(4.85,'COMPARE','원두 특성','품종·산지·배전·인증을 구분해서 비교한다.'),(9.1,'TEST','품질의 증거','독립 점수와 관능평가가 실제로 있는지 확인한다.')]:card(sl,x,3.25,3.65,3.15,lab,head,body)
note(sl,'가격은 돈으로 비교하기 쉽지만 품질은 단일한 수치가 아닙니다. 품종·산지·배전은 커피의 특성을 말해 주고, 인증도 각기 다른 목적을 가집니다. 비싼 브랜드에만 관찰되는 원두 특성이 있는지 가설을 세웠지만, 어느 쪽이 더 좋다고 미리 정하지 않았습니다. 맛의 우열은 블라인드 테이스팅 없이 결론 내리지 않습니다.')

sl=base(3,'02 / The sample','비교할 여섯 브랜드, 두 가지 시장 포지셔닝','그룹은 연구 설계상의 구분이며, 원두 품질 등급을 뜻하지 않는다.')
for col,(grp,items,fill) in enumerate([('A / 일반 프랜차이즈',['스타벅스','투썸플레이스','폴 바셋'],BROWN),('B / 가성비 포지셔닝',['빽다방','텐퍼센트커피','컴포즈커피'],GREEN)]):
    x=.6+col*6.35;rect(sl,x,3.2,5.8,.55,fill,radius=True);text(sl,grp,x+.2,3.32,5.3,.35,16,CREAM,True)
    for k,b in enumerate(items):text(sl,f'0{k+1}    {b}',x+.22,4.0+k*.76,5.2,.55,24,INK,True);line(sl,x+.22,4.62+k*.76,5.25)
note(sl,'사용자가 지정한 여섯 브랜드를 목적 표집했습니다. A는 일반 프랜차이즈, B는 가성비를 중심으로 묶은 분석상의 포지셔닝입니다. 각 그룹 세 브랜드뿐이므로 한국 전체 카페 시장으로 통계적으로 일반화할 수 없습니다. 동일 브랜드 안에서도 원두 옵션과 매장 운영은 달라질 수 있습니다.')

sl=base(4,'03 / Evidence pipeline',f"{M['source_count']}개 근거 → 매장 블렌드 연결 → HOT 6건",'Firecrawl 본문 수집 후 출처·적용일·메뉴 온도·검증 상태를 사람이 읽어 정제했다.',source='수집 2026.10.18 | data/sources.csv · price_observations.csv · collection_log.csv')
for i,(n,h,b) in enumerate([('01','공식 우선','공식 원두 소개·메뉴·공정 자료'),('02','원문 수집','Firecrawl + 공개 웹·HTTP'),('03','교차 검토','온도·시점·제품 적용 범위 분리'),('04','분석 데이터','미확보와 미공개를 구분')]):
    x=.6+i*3.17;rect(sl,x,3.25,2.95,2.1,WHITE,radius=True);text(sl,n,x+.2,3.42,2,.4,24,ACCENT,True,font='Arial');text(sl,h,x+.2,3.98,2.55,.4,19,INK,True);text(sl,b,x+.2,4.56,2.55,.7,13,MUTED)
for x,label,desc,color in [(0.65,'verified','공식 문구/원문 확인',GREEN),(3.85,'reported','외부 보도·재게시',BROWN),(7.05,'claimed','브랜드의 품질 주장',ACCENT),(10.25,'unknown','미확인·미공개',MUTED)]:text(sl,label,x,5.75,2.7,.35,14,color,True,font='Arial');text(sl,desc,x,6.2,2.8,.35,11,MUTED)
note(sl,'공식·취재 근거 출처는 공식 홈페이지와 제품 자료를 우선하고, 가격은 인상 보도와 매장 메뉴를 보완해 수집했습니다. Firecrawl 연결과 실제 본문을 확인했고, 컴포즈 공식 사이트의 국가 제한 403과 스타벅스 잘못된 주소의 소프트 404는 성공 자료에서 제외했습니다. 온도가 명시되지 않은 텐퍼센트 가격 2건과 100원 이벤트는 HOT 비교에 사용하지 않았습니다. verified는 공개 문구를 확인했다는 뜻이지 독립적으로 원두를 시험했다는 뜻이 아닙니다.','S04|S06|S09|S11|S13|S15|S16')

sl=base(5,'04 / The price gap','관측 가격 차이는 최대 3,200원, 약 3.13배','동일시점 실매장 조사 아님 · 판매량 가중치 없는 공개 가격 관측값',source='S04·S08·S31·S12·S14·S17 | 적용 2025.01–05 및 적용일 미확인 매장 자료 · 수집 2026.10.18')
chart(sl,B['브랜드명'].tolist(),B['기본가격'].tolist(),.45,3.05,8.5,3.6,pointcolors=[BROWN]*3+[GREEN]*3,maxval=5600)
text(sl,'3.13×',9.35,3.2,3.3,.85,44,BROWN,True,font='Georgia');text(sl,'자료상 최고 / 최저',9.4,4.15,3,.4,15,MUTED)
chart(sl,['A 평균','B 평균'],[M['group_a_mean'],M['group_b_mean']],9.15,4.72,3.65,1.45,pointcolors=[BROWN,GREEN],maxval=6200)
text(sl,'각 3개 관측값의 단순 평균\n전국 시장 평균을 의미하지 않음',9.4,6.22,3.1,.65,10,MUTED)
note(sl,'공개 자료의 HOT 관측값은 스타벅스 톨 4700원, 투썸 레귤러와 폴 바셋 스탠더드 4700원, 빽다방 1700원, 텐퍼센트 울산비즈파크점 2500원, 컴포즈 1500원입니다. 자료상 최고와 최저는 3200원 차이, 약 3.13배입니다. A의 단순 평균은 4700원, B는 1900원입니다. 다만 가격 적용일이 다르고 텐퍼센트는 특정 매장 재게시 가격입니다. 스타벅스 톨은 최소 판매 크기인 숏과 다릅니다. 2026년 전국 최신 동일가라고 주장하지 않습니다. 공식 HOT 용량은 스타벅스 355mL와 폴 바셋 360mL만 확인해 6개 브랜드의 100mL 가격 순위는 만들지 않았습니다.','S04|S05|S08|S30|S31|S32|S12|S14|S17')

sl=base(6,'05 / Recipe disclosure','원두를 쓴다는 사실과 배합 비율은 다르다','품종 비율을 현재 비교 범위에 연결해 확인한 브랜드는 1/6',source='S01·S06·S13·S24–S30 | 파스쿠찌 제외·폴 바셋 대체 · 품종 비율과 스페셜티 함량은 별개')
for i,b in enumerate(B['브랜드명']):
    y=3.16+i*.49;text(sl,b,.65,y,2,.35,14,INK,True)
    rect(sl,2.85,y+.02,5.9,.29,BROWN if i==0 else BEIGE,radius=True)
    text(sl,'아라비카 100%' if i==0 else '이번 조사에서 비율 근거 미확보',3.04,y+.04,5.5,.27,11,CREAM if i==0 else MUTED)
card(sl,9.25,3.1,3.45,3.28,'SAMPLE CHANGE','파스쿠찌 → 폴 바셋','생산국과 매장 메뉴 연결이\n확인되는 유사 가격대 브랜드로\n교체했다. 비율은 추정하지 않음.',fill=WHITE)
note(sl,'파스쿠찌의 현재 한국 매장 골든삭 생산국은 미확보여서 유사한 4,700원대의 폴 바셋으로 교체했습니다. 폴 바셋 공식 HOT 메뉴는 시그니처 블렌드를 사용한다고 명시하고, 공식 원두 안내는 브라질·에티오피아 생산국을 공개합니다. 품종 비율과 국가별 배합비, 스페셜티 구성분 함량은 서로 다른 정보입니다. 스타벅스 외 품종 비율은 이번에 확보하지 못했습니다. 이것을 기업의 비공개 선언으로 해석하지 않았습니다.','S01|S06|S13|S24|S26|S28|S30')

sl=base(7,'06 / Bean profiles','매장 블렌드의 생산국과 농장을 추적했다','원산지 존재·배전 문구를 비교하며, 산지별 비율과 맛의 우열은 추정하지 않는다.',source='S02·S06·S13·S24–S30 | 생산국·농장과 로스팅 장소는 별개 · 세부 목록은 origin_audit.csv')
rows=[['스타벅스','라틴 아메리카·아시아/태평양¹','다크','캐러멜·강한 바디'],['투썸플레이스','블랙: 콜롬비아·브라질·PNG','다크 / 미디엄','초콜릿·견과 / 과일·꽃'],['폴 바셋','브라질·에티오피아','국내 로스팅','초콜릿·산미·바디'],['빽다방','브라질: 벨로소 DOT·파젠다 엄³','5–15일 사용⁴','브라운슈가·견과'],['텐퍼센트커피','다크: 브라질·콜롬비아·과테말라·에티오피아','강배전 / 중배전²','초콜릿 / 과일·산미'],['컴포즈커피','브라질·콜롬비아·에티오피아³','자체 로스팅³','고소함·초콜릿³']]
grid(sl,['브랜드','산지 / 구성','로스팅','향미 문구'],rows,[2,4,2.4,3.6],y=3.02,rowh=.43,size=11)
text(sl,'¹ 스타벅스 본문은 라틴·아시아, 하단 Origin은 아프리카/아라비아 → 국가 구성 확정 보류\n² 텐퍼센트 공식 용어 병기  ³ 2026 취재·기업 인터뷰  ⁴ 로스팅 후 사용 기간, 배전 단계 아님',.67,6.28,11.9,.53,10,MUTED)
note(sl,'생두 생산국과 로스팅 장소를 분리했습니다. 투썸 블랙그라운드는 콜롬비아·브라질·파푸아뉴기니, 아로마노트는 에티오피아·과테말라입니다. 폴 바셋 시그니처는 브라질·에티오피아이며 국내 로스팅입니다. 빽다방은 2026년 6월 현장 취재에서 브라질 벨로소 DOT, 파젠다 엄 내추럴, 파젠다 엄 펄프드 내추럴 3종을 확인했습니다. 5–15일은 로스팅 후 사용 기간이며 배전 단계가 아닙니다. 텐퍼센트 다크는 브라질·콜롬비아·과테말라·에티오피아, 미디엄은 과테말라·콜롬비아·에티오피아입니다. 컴포즈 2026 기업 인터뷰는 비터홀릭의 브라질·콜롬비아·에티오피아 구성을 안내합니다. 스타벅스 한국 제품 본문은 라틴 아메리카·아시아/태평양, 하단은 아프리카/아라비아로 충돌해 국가를 확정하지 않았습니다. 상세 농장·지역·시점·URL은 origin_audit.csv에 있습니다.','S02|S06|S07|S13|S24|S25|S26|S27|S28|S29|S30')

sl=base(8,'07 / Claim versus evidence','‘스페셜티’라는 문구와 독립 평가표를 구분한다','빽다방 20%는 공식 안내, 컴포즈 5%·구성 원두 82점은 2026년 취재 보도다.',source='S03·S07·S13·S24·S26·S28–S30 | 공식 주장과 보도 구분 · 82점은 전체 블렌드에 전용 불가')
grid(sl,['브랜드 / 범위','확인한 내용','해석 범위'],[['스타벅스','C.A.F.E. Practices 안내','윤리 구매 검증 ≠ 관능 점수'],['투썸플레이스','HACCP·내부 Q-Grader 안내','공정·내부 관리 ≠ 독립 맛 평가'],['폴 바셋','스페셜티 선별·3–14일 사용','공식 관리 주장 ≠ 독립 맛 평가'],['빽다방','파젠다 엄 스페셜티 20%','공식 배합 안내·개별 평가표 미확보'],['텐퍼센트커피','상위 10% 스페셜티 주장','10%는 아라비카 함량이 아님'],['컴포즈커피','스페셜티 5%·구성 원두 82점 보도','완성 블렌드 점수 아님·평가표 미확보']],[2.25,4.7,5.0],y=3.02,rowh=.43,size=11.5)
text(sl,'0 / 6',.7,6.22,1.7,.45,25,BROWN,True,font='Georgia');text(sl,'동일 기준으로 비교 가능한 독립 관능평가 점수 확보 건수',2.5,6.31,9.9,.4,14,MUTED)
note(sl,'빽다방 공식 창업 자료는 파젠다 엄 스페셜티 20% 배합을 명시합니다. 문구의 존재는 확인했지만 개별 로트 평가표를 검증한 것은 아닙니다. 컴포즈 2026년 3월 취재 기사에는 브라질 미나스제라이스 스페셜티 구성분 5%와 해당 원두 SCA 82점이 보도됐습니다. 82점은 전체 비터홀릭이나 매장 한 잔의 독립 점수로 사용할 수 없고 평가 원본도 미확보입니다. 폴 바셋의 스페셜티 선별과 로스팅 후 3–14일 사용은 공식 관리 주장입니다. 텐퍼센트 상위 10%는 실제 배합 함량이 아닙니다. 여섯 브랜드를 같은 조건에서 비교한 독립 평가표는 확보하지 못했습니다.','S03|S07|S13|S24|S26|S28|S29|S30')

sl=base(9,'08 / Price & bean information','함량까지 확인한 가성비 브랜드도 있다','공개·보도된 스페셜티 구성 원두 함량만 표시. 미확보 브랜드는 0%가 아니다.',source='S24 공식 빽다방 안내 · S26 아이뉴스24 2026.03.15 | 스페셜티 구성 함량 · 품질 점수 아님')
sl.shapes.add_picture(str(R/'analysis/figures/06_specialty_share.png'),Inches(.5),Inches(3.0),width=Inches(8.4),height=Inches(3.75))
text(sl,'20% / 5%\n≠\n품질 순위',9.35,3.35,3.2,1.85,27,BROWN,True);text(sl,'배합 함량은 구성의 정보다.\n독립 평가와 공급가가 없어\n맛·원가 우열은 판단 불가.',9.35,5.4,3.2,1.0,14,MUTED)
note(sl,'막대는 빽다방 공식 안내 20%와 컴포즈 취재 보도 5%의 스페셜티 구성 원두 함량입니다. 국가별 전체 배합비나 품종 비율, 품질 점수가 아닙니다. 다른 브랜드는 함량 근거 미확보로 표시하지 않았고 0%로 대체하지 않았습니다. 텐퍼센트 상위 10%를 함량으로 사용하지 않았습니다. 폴 바셋도 스페셜티 선별을 주장하므로 일반 그룹이 스페셜티를 사용하지 않는다고 말할 수 없습니다. 함량이 높을수록 맛이나 공급 원가가 우수한지는 검증하지 못했습니다.','S13|S24|S26|S28|S29')

sl=base(10,'09 / The value of a cafe','원두의 가치와 카페 경험의 가치는 다르다','공식 자료에서 확인한 관리·공간·서비스 사례를 분리했다. 원가 비중은 산출하지 않았다.',source='S02·S03·S05·S07·S09·S13·S18·S21 | 개념 다이어그램 · 요소의 면적·순서는 비용 기여도를 뜻하지 않음')
for i,(label,body) in enumerate([('원두','산지·블렌딩·선별\n브랜드별 공개 구성'),('제조·품질 관리','투썸: 자체 플랜트·선별\n컴포즈: 계열사 로스팅 소개'),('공간·브랜드 경험','투썸: 디저트 카페 공간\n텐퍼센트: 공간·경험 설명'),('운영·서비스','스타벅스: 사이렌오더\n교육·추출 관리 안내')]):
    x=.6+i*3.17;rect(sl,x,3.45,2.95,2.8,WHITE,radius=True);text(sl,f'0{i+1}',x+.2,3.63,2.5,.45,27,ACCENT,True,font='Georgia');text(sl,label,x+.2,4.29,2.55,.5,18,INK,True);text(sl,body,x+.2,5.08,2.55,.9,12,MUTED)
note(sl,'판매가격은 한 잔을 제공하는 운영 전체와 연결되어 있습니다. 공식 자료에는 원두와 블렌딩 외에도 선별 설비, 로스팅 플랜트, 내부 관리, 매장 공간, 디저트, 주문 편의가 등장합니다. 이것들은 실제로 확인한 서비스와 관리 요소입니다. 다만 해당 요소 때문에 브랜드 간 가격이 얼마 차이 나는지, 임대료나 마진이 몇 퍼센트인지 산출하지 않았습니다. 가성비 그룹에도 공간과 경험을 말하는 브랜드가 있으므로 A는 공간, B는 테이크아웃이라는 이분법으로 단정하지 않습니다.','S02|S03|S05|S07|S09|S13|S18|S21')

sl=base(11,'10 / The answer','비싸다는 사실만으로 더 좋은 원두라 판정할 수 없다.','이번 공개 자료 분석이 허용하는 결론',dark=True,source='분석 대상 6개 브랜드 · 가격 적용일·매장·제품 범위 상이 · 동일 기준의 독립 품질 측정값 미확보')
for y,label,msg in [(3.30,'OBSERVED','자료상 HOT 가격은 1,500–4,700원으로 다르다.'),(4.33,'OBSERVED','가성비 그룹에도 스페셜티 사용 주장·보도가 있다.'),(5.36,'NOT ESTABLISHED','가격과 맛·원두의 절대 품질 사이의 관계는 검증하지 못했다.')]:
    text(sl,label,.65,y,3,.35,10,'C3AE98',True,font='Arial');text(sl,msg,3.45,y,9.2,.85,22,CREAM,True)
note(sl,'질문에 대한 답은 비싸다는 사실만으로 더 좋은 원두를 사용한다고 판정할 수 없다는 것입니다. 공개 가격 차이는 관찰됐고 가성비 브랜드에도 스페셜티 사용 주장과 관리 특징이 있습니다. 하지만 이것이 싼 원두가 더 좋다는 반대 결론을 뜻하지도 않습니다. 자료 부족 때문에 품질 차이가 없다는 가설을 입증한 것이 아닙니다. 가격과 객관적인 원두 품질의 관계라는 핵심 가설은 미검증으로 남습니다.')

sl=base(12,'11 / What comes next','다음 검증은 같은 날, 같은 조건의 블라인드 테스트','웹 정보의 공백을 실제 측정과 관능평가로 보완한다.')
for i,(lab,head,body) in enumerate([('LIMITATION','비교 조건 불일치','시점·지점·용량이 다르다.\n공급가·투입량을 확보하지 못했다.'),('MEASURE','동일 조건 구매','같은 날 HOT 기본 메뉴 구매.\n영수증·실제 부피·온도를 기록한다.'),('EXPERIMENT','익명 코드 시음','무작위 순서·동일 온도·반복 시음.\n선호와 향미 평가를 분리한다.')]):card(sl,.6+i*4.22,3.25,3.98,3.15,lab,head,body)
note(sl,'후속 연구에서는 같은 날 동일 생활권에서 여섯 브랜드의 HOT 메뉴를 구매하고 영수증과 실제 부피, 온도를 기록하겠습니다. 시음은 브랜드와 가격을 가린 무작위 코드와 순서로 진행하고 같은 온도에서 반복합니다. 개인적 선호 점수와 향미 특성 평가를 구분하고 가능하면 독립 커핑 평가를 추가합니다. 사전 동의한 평가자를 모집해 반복 측정하되, 사람 수가 작다면 여전히 탐색적 결과로 보고합니다. 원두 공급가를 확보하지 못하면 판매가격과 원두 구매 원가의 관계는 계속 미확인입니다.')

sl=base(13,'12 / Sources & reproducibility','출처를 따라가면 같은 분석을 다시 만들 수 있다.','주요 출처 링크 · 전체 출처와 필드별 근거는 sources.csv·facts.csv·발표 노트에 보존',source='수집일 2026.10.18 | 공식 소개 자료 게시일 미표기 다수 | 가격 보도 2025.01–05 · 일부 매장 적용일 미확인')
ids=['S01','S02','S06','S13','S24','S25','S26','S27','S28','S29','S30','S31','S12','S17']
for i,sid in enumerate(ids):
    row=S[S['출처ID']==sid].iloc[0];x=.65+(i//7)*6.4;y=3.03+(i%7)*.49
    sh=text(sl,f"[{sid}] {row['브랜드명']} · {row['출처종류']}",x,y,5.8,.29,11,INK,True)
    sh.text_frame.paragraphs[0].runs[0].hyperlink.address=row['출처URL']
    text(sl,row['출처URL'].replace('https://','').split('?')[0][:70],x,y+.26,5.8,.2,8,MUTED)
note(sl,'주요 공식 페이지와 가격 출처를 클릭할 수 있도록 넣었습니다. 전체 URL, 게시일 또는 미표기 상태, 수집 시각, 검증 상태는 CSV에 있습니다. 수집 원문과 실패 로그도 남겼습니다. 데이터 정제, 차트, Manim 영상, PowerPoint 생성 코드를 함께 제출합니다. 발표자는 가격 적용일이 2025년인 보도와 현재 수집일이 다르다는 점, 매장별 차이와 독립 품질 지표 미확보를 반드시 함께 설명해야 합니다.','|'.join(S['출처ID']))

prs.save(P/'coffee_brand_analysis.pptx')
(P/'speaker_notes.md').write_text('# 발표 노트 — 약 8–10분\n\n'+ '\n\n'.join(f"## {i+1:02}. {slides[i]['title']}\n\n{n}" for i,n in enumerate(notes)),encoding='utf-8')
(P/'slides.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')
print('PowerPoint: 13 editable slides + notes')
