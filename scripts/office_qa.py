"""Open real PowerPoint, add native animations, export all slides and PDF, inspect chart data."""
from pathlib import Path
import json,time
import win32com.client
from pptx import Presentation
from PIL import Image,ImageOps,ImageDraw
R=Path(__file__).resolve().parents[1];P=R/'presentation';Q=R/'analysis/qa';Q.mkdir(parents=True,exist_ok=True)
path=P/'coffee_brand_analysis.pptx'
app=win32com.client.DispatchEx('PowerPoint.Application')
pres=app.Presentations.Open(str(path),ReadOnly=False,Untitled=False,WithWindow=False)
report={'opened_by':'Microsoft PowerPoint COM','slides':pres.Slides.Count,'native_charts':[],'animation_effects':0,'native_chart_animation':[],'exported':[]}
assert pres.Slides.Count==13
try:
    for sl in pres.Slides:
        sl.SlideShowTransition.AdvanceOnClick=True
        # Fade transitions, no auto advance during a live presentation.
        try:sl.SlideShowTransition.EntryEffect=1793
        except Exception:pass
        for shape in sl.Shapes:
            if shape.HasChart:
                series=shape.Chart.SeriesCollection(1)
                report['native_charts'].append({'slide':sl.SlideIndex,'categories':list(series.XValues),'values':list(series.Values)})
                try:
                    fx=sl.TimeLine.MainSequence.AddEffect(shape,22,0,3);fx.Timing.Duration=1.1
                    report['native_chart_animation'].append({'slide':sl.SlideIndex,'by':'whole chart wipe (modern Office chart)','success':True})
                except Exception as ex:report['native_chart_animation'].append({'slide':sl.SlideIndex,'success':False,'error':str(ex)[:160]})
            elif sl.SlideIndex in [2,3,4,6,7,8,9,10,11,12] and shape.Top>210 and shape.Top<475:
                # Staggered fades preserve editable shapes and text.
                try:
                    fx=sl.TimeLine.MainSequence.AddEffect(shape,10,0,3);fx.Timing.Duration=.25
                    fx.Timing.TriggerDelayTime=.025
                    report['animation_effects']+=1
                except Exception:pass
    pres.Save()
    pres.SaveAs(str(P/'coffee_brand_analysis.pdf'),32)
    for i in range(1,14):
        out=Q/f'slide_{i:02}.png';pres.Slides(i).Export(str(out),'PNG',1600,900);report['exported'].append(out.name)
    pres.Close()
finally:app.Quit()
# Reopen once more after save to verify the delivered file is not corrupt.
app=win32com.client.DispatchEx('PowerPoint.Application');p=app.Presentations.Open(str(path),ReadOnly=True,Untitled=False,WithWindow=False);report['reopened_slides']=p.Slides.Count;p.Close();app.Quit()
expected=[4700,4700,4700,1700,2500,1500]
assert report['native_charts'][0]['values']==expected,report['native_charts']
deck=Presentation(path);bounds=[]
for i,sl in enumerate(deck.slides,1):
    for sh in sl.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>deck.slide_width+9144 or sh.top+sh.height>deck.slide_height+9144:bounds.append({'slide':i,'shape':sh.name})
report['shape_bounds_violations']=bounds;assert not bounds,bounds
sheet=Image.new('RGB',(1200,4*225),'#e7dfd2')
for i,f in enumerate(sorted(Q.glob('slide_*.png'))):
    im=Image.open(f);im.thumbnail((390,219));sheet.paste(im,((i%3)*400,(i//3)*225))
# Five rows are needed for 13 slides.
sheet=Image.new('RGB',(1200,5*225),'#e7dfd2')
for i,f in enumerate(sorted(Q.glob('slide_*.png'))):
    im=Image.open(f);im.thumbnail((390,219));sheet.paste(im,((i%3)*400,(i//3)*225))
sheet.save(Q/'contact_sheet.jpg',quality=92)
(Q/'powerpoint_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='exported'},ensure_ascii=False,indent=2))
